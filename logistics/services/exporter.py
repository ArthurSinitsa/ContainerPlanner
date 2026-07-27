"""
Экспорт результата расчёта (раскладки товаров по контейнерам) в .xlsx.

Публичный API модуля:
    * :class:`PackingExcelExporter` — построение книги для одной заявки;
    * :class:`PackingExportError`   — экспортировать нечего (нет результатов).

Пример использования::

    exporter = PackingExcelExporter(calc_request)
    stream = exporter.build()      # BytesIO с готовой книгой
    name = exporter.filename       # 'raskladka_29.xlsx'

Данные берутся из сохранённых `PackingResult`, а не из «сырого» результата
PackingService.calculate(): расчёт выполняется в Celery-воркере, а отдаёт файл
web-процесс — общей файловой системы у них может не быть.

Всё, что не перечислено выше, — детали реализации (префикс ``_``).
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from logistics.models import CalculationRequest, PackingResult, Product


class PackingExportError(Exception):
    """Заявку невозможно выгрузить (нет рассчитанных контейнеров)."""


@dataclass(frozen=True)
class _ProductRow:
    """Строка товара внутри одного контейнера."""
    product_id: int
    name: str
    sku: str
    quantity: int


@dataclass(frozen=True)
class _ContainerBlock:
    """Один упакованный контейнер: метрики + список товаров в нём."""
    number: int
    type_name: str
    volume_percent: float
    volume_m3: float
    weight_kg: float
    products: list[_ProductRow]

    @property
    def total_units(self) -> int:
        return sum(p.quantity for p in self.products)


class PackingExcelExporter:
    """
    Строит .xlsx с раскладкой заявки по контейнерам.

    Книга состоит из двух листов:
      * «Сводка»   — по строке на контейнер (заполнение, объём, вес, кол-во);
      * «Раскладка» — блоки по контейнерам со списком товаров (ID, Name, SKU, Qty).
    """

    # ── Оформление ────────────────────────────────────────────────────────
    _DARK = "FF23262E"
    _ACCENT = "FFB18CFF"
    _ZEBRA = "FFF5F5F7"
    _MUTED = "FF6B7280"

    _FMT_PERCENT = '0.0"%"'
    _FMT_VOLUME = "0.00"
    _FMT_WEIGHT = "#,##0.0"
    _FMT_INT = "#,##0"

    _SUMMARY_COLUMNS = (
        ("№", 8),
        ("Тип контейнера", 20),
        ("Заполнение по объёму", 22),
        ("Объём груза, м³", 17),
        ("Вес груза, кг", 15),
        ("Позиций", 11),
        ("Всего единиц", 15),
    )
    _LAYOUT_COLUMNS = (
        ("ID", 12),
        ("Наименование", 56),
        ("SKU", 24),
        ("Кол-во", 12),
    )

    def __init__(self, calc_request: CalculationRequest) -> None:
        self._request = calc_request

    # ── Публичный API ─────────────────────────────────────────────────────

    @property
    def filename(self) -> str:
        return f"raskladka_{self._request.id}.xlsx"

    def build(self) -> BytesIO:
        """
        Собирает книгу и возвращает её как поток в памяти.

        :raises PackingExportError: если у заявки нет рассчитанных контейнеров.
        """
        blocks = self._collect()

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Сводка"
        self._write_summary(summary_sheet, blocks)
        self._write_layout(workbook.create_sheet("Раскладка"), blocks)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    # ── Сбор данных ───────────────────────────────────────────────────────

    def _collect(self) -> list[_ContainerBlock]:
        results = list(
            PackingResult.objects
            .filter(calculation_request=self._request)
            .select_related("container_type")
            .order_by("container_number")
        )
        if not results:
            raise PackingExportError(
                f"У заявки #{self._request.id} нет рассчитанных контейнеров."
            )

        catalog = self._product_catalog(results)
        return [self._build_block(result, catalog) for result in results]

    @staticmethod
    def _product_catalog(results: list[PackingResult]) -> dict[int, Product]:
        """Справочник товаров, встречающихся в раскладке (одним запросом)."""
        product_ids = {
            item.get("product_id")
            for result in results
            for item in (result.products or [])
            if item.get("product_id") is not None
        }
        if not product_ids:
            return {}
        return {
            product.product_id: product
            for product in Product.objects.filter(product_id__in=product_ids)
        }

    @classmethod
    def _build_block(cls, result: PackingResult, catalog: dict[int, Product]) -> _ContainerBlock:
        rows: list[_ProductRow] = []
        for item in result.products or []:
            product_id = item.get("product_id")
            product = catalog.get(product_id)
            rows.append(_ProductRow(
                product_id=product_id if product_id is not None else 0,
                # Имя из базы свежее, чем сохранённое в JSON на момент расчёта
                name=(product.name if product and product.name else item.get("product_name") or "—"),
                sku=(product.sku if product and product.sku else "—"),
                quantity=int(item.get("quantity") or 0),
            ))
        rows.sort(key=lambda row: row.product_id)

        return _ContainerBlock(
            number=result.container_number,
            type_name=result.container_type.name,
            volume_percent=float(result.volume_utilization_percent or 0),
            volume_m3=float(result.total_volume_m3 or 0),
            weight_kg=float(result.total_weight_kg or 0),
            products=rows,
        )

    # ── Лист «Сводка» ─────────────────────────────────────────────────────

    def _write_summary(self, sheet: Worksheet, blocks: list[_ContainerBlock]) -> None:
        self._apply_widths(sheet, self._SUMMARY_COLUMNS)
        row = self._write_title(sheet, sheet_width=len(self._SUMMARY_COLUMNS), blocks=blocks)

        row = self._write_table_header(sheet, row, self._SUMMARY_COLUMNS)
        sheet.freeze_panes = f"A{row}"

        first_data_row = row
        for index, block in enumerate(blocks):
            values = (
                block.number,
                block.type_name,
                block.volume_percent,
                block.volume_m3,
                block.weight_kg,
                len(block.products),
                block.total_units,
            )
            formats = (None, None, self._FMT_PERCENT, self._FMT_VOLUME,
                       self._FMT_WEIGHT, self._FMT_INT, self._FMT_INT)
            self._write_data_row(sheet, row, values, formats, striped=index % 2 == 1)
            row += 1

        total_units = sum(block.total_units for block in blocks)
        total_positions = len({p.product_id for block in blocks for p in block.products})
        average_fill = (
            sum(block.volume_percent for block in blocks) / len(blocks) if blocks else 0
        )
        totals = (
            "ИТОГО",
            f"{len(blocks)} конт.",
            average_fill,
            sum(block.volume_m3 for block in blocks),
            sum(block.weight_kg for block in blocks),
            total_positions,
            total_units,
        )
        formats = (None, None, self._FMT_PERCENT, self._FMT_VOLUME,
                   self._FMT_WEIGHT, self._FMT_INT, self._FMT_INT)
        self._write_data_row(sheet, row, totals, formats, striped=False, bold=True, top_border=True)

        note_row = row + 2
        sheet.cell(row=note_row, column=1,
                   value="Заполнение по объёму — доля занятого габаритного объёма контейнера. "
                         "«Позиций» — количество уникальных товаров.")
        sheet.cell(row=note_row, column=1).font = Font(size=9, italic=True, color=self._MUTED)

        if first_data_row <= row - 1:
            sheet.auto_filter.ref = (
                f"A{first_data_row - 1}:"
                f"{get_column_letter(len(self._SUMMARY_COLUMNS))}{row - 1}"
            )

    # ── Лист «Раскладка» ──────────────────────────────────────────────────

    def _write_layout(self, sheet: Worksheet, blocks: list[_ContainerBlock]) -> None:
        self._apply_widths(sheet, self._LAYOUT_COLUMNS)
        width = len(self._LAYOUT_COLUMNS)
        row = self._write_title(sheet, sheet_width=width, blocks=blocks)

        for block in blocks:
            row = self._write_block(sheet, row, block, width)
            row += 1  # пустая строка между контейнерами

    def _write_block(self, sheet: Worksheet, row: int, block: _ContainerBlock, width: int) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        head = sheet.cell(row=row, column=1, value=f"КОНТЕЙНЕР {block.number} · {block.type_name}")
        head.font = Font(bold=True, size=12, color=self._DARK)
        head.fill = PatternFill("solid", start_color=self._ACCENT)
        head.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        sheet.row_dimensions[row].height = 22
        row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        meta = sheet.cell(
            row=row, column=1,
            value=(
                f"Заполнение по объёму: {block.volume_percent:.1f}%    ·    "
                f"Объём груза: {block.volume_m3:.2f} м³    ·    "
                f"Вес груза: {block.weight_kg:,.1f} кг".replace(",", " ")
            ),
        )
        meta.font = Font(size=10, italic=True, color=self._MUTED)
        meta.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        row += 1

        row = self._write_table_header(sheet, row, self._LAYOUT_COLUMNS)

        for index, product in enumerate(block.products):
            self._write_data_row(
                sheet, row,
                (product.product_id, product.name, product.sku, product.quantity),
                (None, None, None, self._FMT_INT),
                striped=index % 2 == 1,
            )
            row += 1

        self._write_data_row(
            sheet, row,
            ("", "Итого по контейнеру", "", block.total_units),
            (None, None, None, self._FMT_INT),
            striped=False, bold=True, top_border=True,
        )
        return row + 1

    # ── Низкоуровневые помощники ──────────────────────────────────────────

    def _write_title(self, sheet: Worksheet, sheet_width: int, blocks: list[_ContainerBlock]) -> int:
        title = sheet.cell(row=1, column=1, value=f"Раскладка по контейнерам · Заявка #{self._request.id}")
        title.font = Font(bold=True, size=14, color=self._DARK)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sheet_width)

        created = self._request.created_at
        parts = [f"Создана: {created:%d.%m.%Y %H:%M}"] if created else []
        if self._request.description:
            parts.append(str(self._request.description))
        parts.append(f"Контейнеров: {len(blocks)}")

        subtitle = sheet.cell(row=2, column=1, value="    ·    ".join(parts))
        subtitle.font = Font(size=10, color=self._MUTED)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=sheet_width)

        return 4  # строка 3 остаётся пустой

    def _write_table_header(self, sheet: Worksheet, row: int,
                            columns: tuple[tuple[str, int], ...]) -> int:
        for index, (caption, _) in enumerate(columns, start=1):
            cell = sheet.cell(row=row, column=index, value=caption)
            cell.font = Font(bold=True, size=10, color="FFFFFFFF")
            cell.fill = PatternFill("solid", start_color=self._DARK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self._border()
        sheet.row_dimensions[row].height = 26
        return row + 1

    def _write_data_row(self, sheet: Worksheet, row: int, values: tuple,
                        formats: tuple, striped: bool,
                        bold: bool = False, top_border: bool = False) -> None:
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            number_format = formats[index - 1]
            if number_format:
                cell.number_format = number_format
            cell.font = Font(size=10, bold=bold)
            cell.alignment = Alignment(
                horizontal="left" if index == 2 else "center",
                vertical="center",
            )
            cell.border = self._border(top=top_border)
            if striped:
                cell.fill = PatternFill("solid", start_color=self._ZEBRA)

    @staticmethod
    def _border(top: bool = False) -> Border:
        thin = Side(style="thin", color="FFDDDDE3")
        medium = Side(style="medium", color="FF9AA0A6")
        return Border(left=thin, right=thin, bottom=thin, top=medium if top else thin)

    @staticmethod
    def _apply_widths(sheet: Worksheet, columns: tuple[tuple[str, int], ...]) -> None:
        for index, (_, width) in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
