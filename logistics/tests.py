from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from .models import CalculationRequest, ContainerType, PackingResult, Product
from .services.exporter import PackingExcelExporter, PackingExportError

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ExportTestDataMixin:
    """Общая фикстура: заявка с двумя рассчитанными контейнерами."""

    def setUp(self):
        self.container_type = ContainerType.objects.create(
            name="40HQ", length_mm=12000, width_mm=2350, height_mm=2690,
            max_weight_kg=26000, volume_m3=76.0,
        )
        Product.objects.create(product_id=1024, sku="TV-55", name="Телевизор 55", category="TV")
        Product.objects.create(product_id=2048, sku="MON-27", name="Монитор 27", category="Monitor")

        self.calc_request = CalculationRequest.objects.create(status='COMPLETED', description="Тестовая заявка")
        PackingResult.objects.create(
            calculation_request=self.calc_request, container_number=1, container_type=self.container_type,
            total_weight_kg=6120.5, total_volume_m3=65.32,
            volume_utilization_percent=87.4, area_utilization_percent=95.1,
            products=[
                {"product_id": 1024, "product_name": "Устаревшее имя", "quantity": 120},
                {"product_id": 2048, "product_name": "Монитор 27", "quantity": 480},
            ],
            packing_layout=[],
        )
        PackingResult.objects.create(
            calculation_request=self.calc_request, container_number=2, container_type=self.container_type,
            total_weight_kg=2980.0, total_volume_m3=31.5,
            volume_utilization_percent=42.0, area_utilization_percent=61.3,
            products=[{"product_id": 1024, "product_name": "Телевизор 55", "quantity": 60}],
            packing_layout=[],
        )


class PackingExcelExporterTests(ExportTestDataMixin, TestCase):

    def _workbook(self, calc_request=None):
        stream = PackingExcelExporter(calc_request or self.calc_request).build()
        self.assertIsInstance(stream, BytesIO)
        return load_workbook(stream)

    def test_filename_contains_request_id(self):
        exporter = PackingExcelExporter(self.calc_request)
        self.assertEqual(exporter.filename, f"raskladka_{self.calc_request.id}.xlsx")

    def test_workbook_has_two_sheets(self):
        self.assertEqual(self._workbook().sheetnames, ["Сводка", "Раскладка"])

    def test_summary_row_per_container(self):
        sheet = self._workbook()["Сводка"]
        rows = [r for r in sheet.iter_rows(min_row=5, values_only=True) if r[0] in (1, 2)]
        self.assertEqual(len(rows), 2)

        first = rows[0]
        self.assertEqual(first[0], 1)                # № контейнера
        self.assertEqual(first[1], "40HQ")           # тип
        self.assertAlmostEqual(first[2], 87.4)       # заполнение по объёму
        self.assertAlmostEqual(first[3], 65.32)      # объём груза
        self.assertAlmostEqual(first[4], 6120.5)     # вес груза
        self.assertEqual(first[5], 2)                # позиций
        self.assertEqual(first[6], 600)              # всего единиц

    def test_summary_totals(self):
        sheet = self._workbook()["Сводка"]
        totals = next(r for r in sheet.iter_rows(min_row=5, values_only=True) if r[0] == "ИТОГО")
        self.assertAlmostEqual(totals[3], 96.82)                 # суммарный объём
        self.assertAlmostEqual(totals[4], 9100.5)                # суммарный вес
        self.assertEqual(totals[5], 2)                           # уникальных товаров во всей заявке
        self.assertEqual(totals[6], 660)                         # всего единиц

    def test_layout_contains_sku_from_catalog(self):
        sheet = self._workbook()["Раскладка"]
        rows = [r for r in sheet.iter_rows(values_only=True) if r[0] in (1024, 2048)]
        self.assertEqual(len(rows), 3)  # 1024 встречается в обоих контейнерах

        by_id = {}
        for row in rows:
            by_id.setdefault(row[0], []).append(row)

        # SKU подтягивается из справочника — в JSON результата его нет
        self.assertEqual(by_id[1024][0][2], "TV-55")
        self.assertEqual(by_id[2048][0][2], "MON-27")
        # Имя берётся из актуальной карточки товара, а не из «замороженного» JSON
        self.assertEqual(by_id[1024][0][1], "Телевизор 55")
        # Количество — именно в этом контейнере
        self.assertEqual([row[3] for row in by_id[1024]], [120, 60])

    def test_unknown_product_falls_back_to_stored_name(self):
        result = PackingResult.objects.get(calculation_request=self.calc_request, container_number=1)
        result.products = [{"product_id": 9999, "product_name": "Товар вне справочника", "quantity": 7}]
        result.save()

        sheet = self._workbook()["Раскладка"]
        row = next(r for r in sheet.iter_rows(values_only=True) if r[0] == 9999)
        self.assertEqual(row[1], "Товар вне справочника")
        self.assertEqual(row[2], "—")
        self.assertEqual(row[3], 7)

    def test_request_without_results_raises(self):
        empty = CalculationRequest.objects.create(status='COMPLETED')
        with self.assertRaises(PackingExportError):
            PackingExcelExporter(empty).build()


class ExportEndpointTests(ExportTestDataMixin, TestCase):

    def test_template_download(self):
        response = self.client.get('/api/calculate/template/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], XLSX_CONTENT_TYPE)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('shablon_zayavki.xlsx', response['Content-Disposition'])

        content = b"".join(response.streaming_content)
        self.assertTrue(content.startswith(b'PK'))  # zip-контейнер xlsx

    def test_export_returns_xlsx(self):
        response = self.client.get(f'/api/calculate/{self.calc_request.id}/export/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], XLSX_CONTENT_TYPE)
        self.assertIn(f'raskladka_{self.calc_request.id}.xlsx', response['Content-Disposition'])

        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(workbook.sheetnames, ["Сводка", "Раскладка"])

    def test_export_unknown_request_returns_404(self):
        response = self.client.get('/api/calculate/999999/export/')
        self.assertEqual(response.status_code, 404)
        self.assertIn('error', response.json())

    def test_export_unfinished_request_returns_409(self):
        self.calc_request.status = 'PROCESSING'
        self.calc_request.save()

        response = self.client.get(f'/api/calculate/{self.calc_request.id}/export/')
        self.assertEqual(response.status_code, 409)
        self.assertIn('error', response.json())

    def test_export_completed_without_results_returns_409(self):
        empty = CalculationRequest.objects.create(status='COMPLETED')
        response = self.client.get(f'/api/calculate/{empty.id}/export/')
        self.assertEqual(response.status_code, 409)
        self.assertIn('error', response.json())
